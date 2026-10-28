import io
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from django.db.models import Count, DecimalField, Q, Sum, Value
from django.db.models.functions import Coalesce, TruncMonth
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.response import Response
from rest_framework.views import APIView
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from xml.sax.saxutils import escape
from apps.property.enums import Category, TransactionType
from apps.property.models import Property, Finance
from apps.organisation.utils import get_request_organisation
from common.permission import (
    IsLandlord,
    IsMortgageAdviser,
    IsAdmin,
    IsSuperAdmin,
)


class PropertyPortfolioExportView(APIView):
    permission_classes = [IsLandlord | IsAdmin | IsMortgageAdviser | IsSuperAdmin ]

    HEADERS = [
        "Property Address",
        "Property Type",
        "Number Of Bedrooms",
        "Owners",
        "Original Purchase Date",
        "Original Purchase Price",
        "Current Market Value",
        "Mortgage Lender Name",
        "Outstanding Balance",
        "Current Interest Rate",
        "Monthly Mortgage Payment",
        "Repayment Method",
        "Mortgage End Date",
        "Interest Rate Expiry Date",
        "Monthly Rental Income",
        "EPC Rating",
        "Property Tenure",
        "If Leasehold, Remaining Lease Term",
        "Monthly Service Charge",
    ]

    def get(self, request):
        export_format = request.query_params.get("export_format", "xlsx").lower()
        if export_format not in ("xlsx", "pdf"):
            return HttpResponse(
                "Invalid format. Use ?export_format=xlsx or ?export_format=pdf.",
                status=400,
            )

        organisation = get_request_organisation(request)
        if not organisation:
            return HttpResponse("You are not part of any organisation.", status=400)

        properties = (
            Property.objects.filter(organisation=organisation)
            .prefetch_related("property_mortgages", "shareholder")
            .order_by("-created_at")
        )

        rows = self._build_rows(properties)
        stamp = date.today().isoformat()

        if export_format == "xlsx":
            buffer = self._build_xlsx(rows)
            filename = f"Property_Portfolio_Summary_-_Landkeeper_{stamp}.xlsx"
            content_type = (
                "application/vnd.openxmlformats-officedocument" ".spreadsheetml.sheet"
            )
        else:
            buffer = self._build_pdf(rows)
            filename = f"Property_Portfolio_Summary_-_Landkeeper_{stamp}.pdf"
            content_type = "application/pdf"

        response = HttpResponse(buffer.getvalue(), content_type=content_type)
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    # Data assembly
    def _build_rows(self, properties):
        rows = []
        for prop in properties:
            mortgage = prop.property_mortgages.order_by("-created_at").first()
            owners = ", ".join(
                filter(None, (o.owner_name for o in prop.shareholder.all()))
            )

            rows.append(
                {
                    "address": prop.address,
                    "property_type": prop.get_property_type_display(),
                    "bedrooms": prop.bedrooms,
                    "owners": owners,
                    "purchase_date": prop.purchase_date,
                    "purchase_price": prop.purchase_price,
                    "current_value": prop.current_value,
                    "lender": mortgage.lender_name if mortgage else "",
                    "outstanding_balance": (
                        mortgage.outstanding_balance if mortgage else None
                    ),
                    "interest_rate": mortgage.interest_rate if mortgage else None,
                    "monthly_payment": mortgage.monthly_payment if mortgage else None,
                    "repayment_method": (
                        mortgage.get_interest_rate_type_display() if mortgage else ""
                    ),
                    "mortgage_end_date": None,
                    "rate_expiry_date": (
                        mortgage.interest_rate_expiry_date if mortgage else None
                    ),
                    "monthly_rental_income": prop.monthly_rental_income,
                    "epc": mortgage.epc_rating if mortgage else "",
                    "tenure": (
                        prop.get_property_tenure_display()
                        if prop.property_tenure
                        else ""
                    ),
                    "lease_term": prop.remaining_lease_term,
                    "service_charge": prop.monthly_service_charge,
                }
            )
        return rows

    def _totals(self, rows):
        def total(key):
            return sum((r[key] or 0) for r in rows)

        current_value = total("current_value")
        outstanding_balance = total("outstanding_balance")
        monthly_payment = total("monthly_payment")
        monthly_rental_income = total("monthly_rental_income")
        return {
            "current_value": current_value,
            "outstanding_balance": outstanding_balance,
            "monthly_payment": monthly_payment,
            "monthly_rental_income": monthly_rental_income,
            "net_asset_value": current_value - outstanding_balance,
            "net_income": monthly_rental_income - monthly_payment,
        }

    # Excel export (matches template)
    def _build_xlsx(self, rows):
        LIGHT_BLUE = "DCE6F1"
        GREY = "D9D9D9"

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        header_font = Font(name="Calibri", size=10, bold=True, color="000000")
        header_fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        body_font = Font(name="Calibri", size=10)
        thin = Side(style="thin", color=GREY)
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, label in enumerate(self.HEADERS, start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
        ws.row_dimensions[1].height = 42

        widths = [
            30,
            14,
            10,
            18,
            14,
            15,
            14,
            18,
            15,
            14,
            16,
            16,
            14,
            16,
            15,
            10,
            14,
            18,
            14,
        ]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        currency_cols = {6, 7, 9, 11, 15, 19}
        percent_cols = {10}
        date_cols = {5, 13, 14}

        row_idx = 2
        for r in rows:
            values = [
                r["address"],
                r["property_type"],
                r["bedrooms"],
                r["owners"],
                r["purchase_date"],
                r["purchase_price"],
                r["current_value"],
                r["lender"],
                r["outstanding_balance"],
                r["interest_rate"],
                r["monthly_payment"],
                r["repayment_method"],
                r["mortgage_end_date"],
                r["rate_expiry_date"],
                r["monthly_rental_income"],
                r["epc"],
                r["tenure"],
                r["lease_term"],
                r["service_charge"],
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(
                    horizontal="left" if col == 1 else "center", vertical="center"
                )
                if col in currency_cols:
                    cell.number_format = "£#,##0.00;(£#,##0.00);-"
                elif col in percent_cols:
                    cell.number_format = "0.00%"
                elif col in date_cols:
                    cell.number_format = "dd/mm/yyyy"
            row_idx += 1

        last_data_row = row_idx - 1

        # TOTALS label row (with sub-labels: VALUES / MORTGAGE SUM / PAYMENTS / RENT)
        label_row = row_idx
        ws.cell(row=label_row, column=1, value="TOTALS").font = Font(bold=True)
        ws.cell(row=label_row, column=7, value="VALUES").font = Font(bold=True)
        ws.cell(row=label_row, column=9, value="MORTGAGE SUM").font = Font(bold=True)
        ws.cell(row=label_row, column=11, value="PAYMENTS").font = Font(bold=True)
        ws.cell(row=label_row, column=15, value="RENT").font = Font(bold=True)
        for col in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=label_row, column=col)
            cell.fill = PatternFill("solid", fgColor=GREY)
            cell.border = border

        # Totals values row
        values_row = label_row + 1
        ws.cell(row=values_row, column=7, value=f"=SUM(G2:G{last_data_row})")
        ws.cell(row=values_row, column=9, value=f"=SUM(I2:I{last_data_row})")
        ws.cell(row=values_row, column=11, value=f"=SUM(K2:K{last_data_row})")
        ws.cell(row=values_row, column=15, value=f"=SUM(O2:O{last_data_row})")
        for col in (7, 9, 11, 15):
            cell = ws.cell(row=values_row, column=col)
            cell.font = Font(bold=True)
            cell.number_format = "£#,##0.00"
            cell.fill = PatternFill("solid", fgColor=GREY)
            cell.border = border

        # NET ASSET VALUE / NET INCOME rows
        nav_row = values_row + 1
        income_row = values_row + 2
        ws.cell(row=nav_row, column=1, value="NET ASSET VALUE")
        ws.cell(row=nav_row, column=2, value=f"=G{values_row}-I{values_row}")
        ws.cell(row=income_row, column=1, value="NET INCOME")
        ws.cell(row=income_row, column=2, value=f"=O{values_row}-K{values_row}")
        for r in (nav_row, income_row):
            for col in (1, 2):
                cell = ws.cell(row=r, column=col)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
                if col == 2:
                    cell.number_format = "£#,##0.00"

        generated_row = income_row + 2
        ws.cell(row=generated_row, column=1, value="Report generated on").font = Font(bold=True)
        gen_cell = ws.cell(
            row=generated_row,
            column=2,
            value=timezone.localtime().strftime("%d-%m-%Y %H:%M"),
        )
        gen_cell.font = Font(bold=True)
        gen_cell.alignment = Alignment(horizontal="left")

        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # PDF export
    def _build_pdf(self, rows):
        INK = colors.HexColor("#111827")
        SUBTLE = colors.HexColor("#6B7280")
        HAIRLINE = colors.HexColor("#E5E7EB")
        ACCENT = colors.HexColor("#8A6D3B")
        WHITE = colors.white
        ZEBRA = colors.HexColor("#F9FAFB")

        generated_on = timezone.localtime().strftime("%d-%m-%Y %H:%M")

        page_size = landscape(A4)
        page_w, _ = page_size

        title_style = ParagraphStyle(
            "Title",
            fontName="Helvetica-Bold",
            fontSize=20,
            textColor=INK,
            alignment=TA_LEFT,
            leading=24,
        )
        subtitle_style = ParagraphStyle(
            "Subtitle",
            fontName="Helvetica",
            fontSize=10,
            textColor=SUBTLE,
        )
        kpi_label_style = ParagraphStyle(
            "KPILabel",
            fontName="Helvetica",
            fontSize=9,
            textColor=SUBTLE,
        )
        kpi_value_style = ParagraphStyle(
            "KPIValue",
            fontName="Helvetica-Bold",
            fontSize=17,
            textColor=INK,
            spaceBefore=2,
        )
        header_left = ParagraphStyle(
            "HeaderLeft",
            fontName="Helvetica-Bold",
            fontSize=6.6,
            textColor=INK,
            leading=8,
        )
        header_center = ParagraphStyle(
            "HeaderCenter",
            fontName="Helvetica-Bold",
            fontSize=6.6,
            textColor=INK,
            alignment=1,
            leading=8,
        )
        # Wrapping body styles so long text doesn't get clipped/overflow the cell
        body_left = ParagraphStyle(
            "BodyLeft",
            fontName="Helvetica",
            fontSize=7.6,
            textColor=INK,
            alignment=TA_LEFT,
            leading=9,
        )
        body_center = ParagraphStyle(
            "BodyCenter",
            fontName="Helvetica",
            fontSize=7.6,
            textColor=INK,
            alignment=1,
            leading=9,
        )

        def fmt_money(v):
            return f"£{v:,.0f}" if v is not None else ""

        def fmt_pct(v):
            return f"{v:.2f}%" if v is not None else ""

        def fmt_date(v):
            return v.strftime("%d/%m/%Y") if v else ""

        # Columns whose text can run long and need wrapping rather than truncation
        wrap_left_cols = {
            0,
            3,
            7,
            11,
            16,
        }  # address, owners, lender, repayment method, tenure

        table_data = [
            [
                Paragraph(label, header_left if i == 0 else header_center)
                for i, label in enumerate(self.HEADERS)
            ]
        ]
        for r in rows:
            raw_values = [
                r["address"] or "",
                r["property_type"] or "",
                str(r["bedrooms"]) if r["bedrooms"] is not None else "",
                r["owners"] or "",
                fmt_date(r["purchase_date"]),
                fmt_money(r["purchase_price"]),
                fmt_money(r["current_value"]),
                r["lender"] or "",
                fmt_money(r["outstanding_balance"]),
                fmt_pct(r["interest_rate"]),
                fmt_money(r["monthly_payment"]),
                r["repayment_method"] or "",
                fmt_date(r["mortgage_end_date"]),
                fmt_date(r["rate_expiry_date"]),
                fmt_money(r["monthly_rental_income"]),
                r["epc"] or "",
                r["tenure"] or "",
                str(r["lease_term"]) if r["lease_term"] is not None else "",
                fmt_money(r["service_charge"]),
            ]
            row_cells = []
            for i, value in enumerate(raw_values):
                if i in wrap_left_cols:
                    row_cells.append(Paragraph(value, body_left))
                else:
                    row_cells.append(Paragraph(value, body_center))
            table_data.append(row_cells)

        totals = self._totals(rows)
        totals_row = [
            "Totals",
            "",
            "",
            "",
            "",
            "",
            fmt_money(totals["current_value"]),
            "",
            fmt_money(totals["outstanding_balance"]),
            "",
            fmt_money(totals["monthly_payment"]),
            "",
            "",
            "",
            fmt_money(totals["monthly_rental_income"]),
            "",
            "",
            "",
            "",
        ]
        table_data.append(totals_row)

        col_widths = [
            33 * mm,
            17 * mm,
            12 * mm,
            14 * mm,
            16 * mm,
            16 * mm,
            17 * mm,
            16 * mm,
            16 * mm,
            13 * mm,
            15 * mm,
            18 * mm,
            15 * mm,
            15 * mm,
            17 * mm,
            11 * mm,
            14 * mm,
            13 * mm,
            15 * mm,
        ]
        scale = (page_w - 20 * mm) / sum(col_widths)
        col_widths = [w * scale for w in col_widths]

        n_body_rows = len(rows)
        table_style_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), WHITE),
            ("VALIGN", (0, 0), (-1, 0), "BOTTOM"),
            ("TOPPADDING", (0, 0), (-1, 0), 4),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("LEFTPADDING", (0, 0), (-1, 0), 3),
            ("RIGHTPADDING", (0, 0), (-1, 0), 3),
            ("LINEBELOW", (0, 0), (-1, 0), 1, INK),
            ("VALIGN", (0, 1), (-1, -2), "MIDDLE"),
            ("TOPPADDING", (0, 1), (-1, -2), 6),
            ("BOTTOMPADDING", (0, 1), (-1, -2), 6),
            ("LEFTPADDING", (0, 1), (-1, -2), 3),
            ("RIGHTPADDING", (0, 1), (-1, -2), 3),
            ("LINEBELOW", (0, 1), (-1, -2), 0.5, HAIRLINE),
            ("LINEABOVE", (0, -1), (-1, -1), 1, INK),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, -1), (-1, -1), 8),
            ("ALIGN", (0, -1), (-1, -1), "CENTER"),
            ("ALIGN", (0, -1), (0, -1), "LEFT"),
            ("TOPPADDING", (0, -1), (-1, -1), 7),
            ("BOTTOMPADDING", (0, -1), (-1, -1), 7),
        ]
        # Zebra striping on body rows for readability
        for i in range(1, n_body_rows + 1):
            if i % 2 == 0:
                table_style_cmds.append(("BACKGROUND", (0, i), (-1, i), ZEBRA))

        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle(table_style_cmds))

        def header_footer(canvas, doc):
            canvas.saveState()
            canvas.setStrokeColor(HAIRLINE)
            canvas.setLineWidth(0.5)
            canvas.line(10 * mm, 12 * mm, page_w - 10 * mm, 12 * mm)
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(SUBTLE)
            canvas.drawString(
                10 * mm, 8 * mm, "Landkeeper  ·  Property Portfolio Summary"
            )
            canvas.drawRightString(page_w - 10 * mm, 8 * mm, f"Page {doc.page}")
            canvas.restoreState()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=page_size,
            leftMargin=10 * mm,
            rightMargin=10 * mm,
            topMargin=14 * mm,
            bottomMargin=16 * mm,
            title="Property Portfolio Summary - Landkeeper",
        )

        story = [
            Paragraph("Property Portfolio Summary", title_style),
            Spacer(1, 3),
            Paragraph("Landkeeper — Landlord Property Management", subtitle_style),
            Spacer(1, 2),
            Paragraph(f"Report generated on {generated_on}", subtitle_style),
            Spacer(1, 14),
            table,
            Spacer(1, 16),
        ]

        kpi_table = Table(
            [
                [
                    Paragraph("NET ASSET VALUE", kpi_label_style),
                    Paragraph("NET INCOME", kpi_label_style),
                ],
                [
                    Paragraph(fmt_money(totals["net_asset_value"]), kpi_value_style),
                    Paragraph(fmt_money(totals["net_income"]), kpi_value_style),
                ],
            ],
            colWidths=[70 * mm, 70 * mm],
        )
        kpi_table.setStyle(
            TableStyle(
                [
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LINEABOVE", (0, 0), (0, 0), 1.2, ACCENT),
                    ("LINEABOVE", (1, 0), (1, 0), 1.2, ACCENT),
                    ("TOPPADDING", (0, 0), (-1, 0), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 0),
                    ("TOPPADDING", (0, 1), (-1, 1), 2),
                ]
            )
        )
        story.append(kpi_table)

        doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
        buffer.seek(0)
        return buffer


class IncomeReportView(APIView):
    """
    Filter: start_date=2026-01-01&end_date=2026-06-30
    """

    permission_classes = [IsLandlord | IsAdmin | IsMortgageAdviser | IsSuperAdmin]

    REPORT_CATEGORIES = [Category.RENTAL_INCOME, Category.OTHER]
    CURRENCY = "£"

    # Formatting helpers
    @classmethod
    def money(cls, value):
        return f"{cls.CURRENCY}{(value or Decimal('0')):,.2f}"

    @staticmethod
    def safe(text):
        return escape(str(text)) if text else "-"

    # Data
    def get_scoped_queryset(self, request):
        qs = Finance.objects.filter(
            type=TransactionType.INCOME,
            category__in=self.REPORT_CATEGORIES,
        ).select_related("property", "organisation")

        if (
            IsAdmin().has_permission(request, self)
            or IsSuperAdmin().has_permission(request, self)
        ):
            return qs

        organisation = request.user.get_organisation()
        if not organisation:
            return qs.none()

        if hasattr(organisation, "pk"):
            return qs.filter(organisation=organisation)
        return qs.filter(organisation__in=organisation)

    def apply_filters(self, request, qs):
        p = request.query_params
        start = p.get("start_date")
        end = p.get("end_date")

        start_d = datetime.strptime(start, "%Y-%m-%d").date() if start else None
        end_d = datetime.strptime(end, "%Y-%m-%d").date() if end else None

        if start_d and end_d and start_d > end_d:
            raise ValueError("'start_date' cannot be after 'end_date'.")

        if start_d:
            qs = qs.filter(date__gte=start_d)
        if end_d:
            qs = qs.filter(date__lte=end_d)

        fmt = "%d %b %Y"
        if start_d and end_d:
            period = f"{start_d:{fmt}} to {end_d:{fmt}}"
        elif start_d:
            period = f"From {start_d:{fmt}}"
        elif end_d:
            period = f"Up to {end_d:{fmt}}"
        else:
            period = "All time"

        return qs, period

    # PDF helpers
    @staticmethod
    def styled_table(data, col_widths):
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F2F2F2")],
                    ),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    # Total row
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#D9E1F2")),
                ]
            )
        )
        return table

    @staticmethod
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, "Income Report")
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    def get(self, request):
        money = self.money
        safe = self.safe

        try:
            qs, period = self.apply_filters(request, self.get_scoped_queryset(request))
        except ValueError as e:
            msg = str(e)
            if "does not match format" in msg:
                msg = "Dates must be in YYYY-MM-DD format."
            return Response({"detail": msg}, status=status.HTTP_400_BAD_REQUEST)

        transactions = qs.order_by("date", "id")
        zero = Value(0, output_field=DecimalField(max_digits=20, decimal_places=2))
        total = qs.aggregate(total=Coalesce(Sum("amount"), zero))["total"]

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=18 * mm,
            title="Income Report",
        )
        styles = getSampleStyleSheet()
        h2 = ParagraphStyle(
            "h2", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79")
        )
        small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9, leading=11)
        w = doc.width
        story = []

        story.append(Paragraph("Income Report", styles["Title"]))
        story.append(Paragraph(f"<b>Period:</b> {safe(period)}", styles["Normal"]))
        story.append(
            Paragraph(
                f"<b>Generated:</b> {timezone.localtime():%d %b %Y, %H:%M}",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 6 * mm))

        story.append(Paragraph("All Transactions", h2))

        if transactions.exists():
            data = [["Date", "Property", "Type", "Category", "Description", "Amount"]]
            for f in transactions:
                data.append(
                    [
                        f.date.strftime("%d/%m/%Y") if f.date else "-",
                        Paragraph(safe(f.property.property_name), small),
                        f.get_type_display(),
                        Paragraph(safe(f.get_category_display()), small),
                        Paragraph(safe(f.description), small),
                        money(f.amount),
                    ]
                )
            data.append(["", "", "", "", "Total", money(total)])

            t = self.styled_table(
                data,
                [w * 0.13, w * 0.20, w * 0.10, w * 0.15, w * 0.25, w * 0.17],
            )
            story.append(t)
        else:
            story.append(
                Paragraph("No transactions found for this period.", styles["Normal"])
            )

        doc.build(story, onFirstPage=self.footer, onLaterPages=self.footer)
        buffer.seek(0)

        filename = f"income_report_{timezone.now():%Y%m%d}.pdf"
        response = HttpResponse(buffer, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class ExpenseReportView(APIView):
    """
    Filter: start_date=2026-01-01&end_date=2026-06-30
    """

    permission_classes = [IsLandlord | IsAdmin | IsMortgageAdviser | IsSuperAdmin]

    REPORT_CATEGORIES = [
        Category.MORTGAGE_PAYMENT,
        Category.REPAIRS,
        Category.INSURANCE,
        Category.SERVICE_CHARGES,
        Category.UTILITIES,
        Category.MANAGEMENT_FEES,
        Category.TAX,
        Category.OTHER,
    ]
    CURRENCY = "£"

    # Formatting helpers
    @classmethod
    def money(cls, value):
        return f"{cls.CURRENCY}{(value or Decimal('0')):,.2f}"

    @staticmethod
    def safe(text):
        return escape(str(text)) if text else "-"

    # Data
    def get_scoped_queryset(self, request):
        qs = Finance.objects.filter(
            type=TransactionType.EXPENSE,
            category__in=self.REPORT_CATEGORIES,
        ).select_related("property", "organisation")

        if (
            IsAdmin().has_permission(request, self)
            or IsSuperAdmin().has_permission(request, self)
        ):
            return qs

        organisation = request.user.get_organisation()
        if not organisation:
            return qs.none()

        if hasattr(organisation, "pk"):
            return qs.filter(organisation=organisation)
        return qs.filter(organisation__in=organisation)

    def apply_filters(self, request, qs):
        p = request.query_params
        start = p.get("start_date")
        end = p.get("end_date")

        start_d = datetime.strptime(start, "%Y-%m-%d").date() if start else None
        end_d = datetime.strptime(end, "%Y-%m-%d").date() if end else None

        if start_d and end_d and start_d > end_d:
            raise ValueError("'start_date' cannot be after 'end_date'.")

        if start_d:
            qs = qs.filter(date__gte=start_d)
        if end_d:
            qs = qs.filter(date__lte=end_d)

        fmt = "%d %b %Y"
        if start_d and end_d:
            period = f"{start_d:{fmt}} to {end_d:{fmt}}"
        elif start_d:
            period = f"From {start_d:{fmt}}"
        elif end_d:
            period = f"Up to {end_d:{fmt}}"
        else:
            period = "All time"

        return qs, period

    # PDF helpers
    @staticmethod
    def styled_table(data, col_widths):
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F2F2F2")],
                    ),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    # Total row
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#D9E1F2")),
                ]
            )
        )
        return table

    @staticmethod
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, "Expense Report")
        canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    def get(self, request):
        money = self.money
        safe = self.safe

        try:
            qs, period = self.apply_filters(request, self.get_scoped_queryset(request))
        except ValueError as e:
            msg = str(e)
            if "does not match format" in msg:
                msg = "Dates must be in YYYY-MM-DD format."
            return Response({"detail": msg}, status=status.HTTP_400_BAD_REQUEST)

        transactions = qs.order_by("date", "id")
        zero = Value(0, output_field=DecimalField(max_digits=20, decimal_places=2))
        total = qs.aggregate(total=Coalesce(Sum("amount"), zero))["total"]

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=18 * mm,
            title="Expense Report",
        )
        styles = getSampleStyleSheet()
        h2 = ParagraphStyle(
            "h2", parent=styles["Heading2"], textColor=colors.HexColor("#1F4E79")
        )
        small = ParagraphStyle("small", parent=styles["Normal"], fontSize=9, leading=11)
        w = doc.width
        story = []

        story.append(Paragraph("Expense Report", styles["Title"]))
        story.append(Paragraph(f"<b>Period:</b> {safe(period)}", styles["Normal"]))
        story.append(
            Paragraph(
                f"<b>Generated:</b> {timezone.localtime():%d %b %Y, %H:%M}",
                styles["Normal"],
            )
        )
        story.append(Spacer(1, 6 * mm))

        story.append(Paragraph("All Transactions", h2))

        if transactions.exists():
            data = [["Date", "Property", "Type", "Category", "Description", "Amount"]]
            for f in transactions:
                data.append(
                    [
                        f.date.strftime("%d/%m/%Y") if f.date else "-",
                        Paragraph(safe(f.property.property_name), small),
                        f.get_type_display(),
                        Paragraph(safe(f.get_category_display()), small),
                        Paragraph(safe(f.description), small),
                        money(f.amount),
                    ]
                )
            data.append(["", "", "", "", "Total", money(total)])

            t = self.styled_table(
                data,
                [w * 0.13, w * 0.20, w * 0.10, w * 0.15, w * 0.25, w * 0.17],
            )
            story.append(t)
        else:
            story.append(
                Paragraph("No transactions found for this period.", styles["Normal"])
            )

        doc.build(story, onFirstPage=self.footer, onLaterPages=self.footer)
        buffer.seek(0)

        filename = f"expense_report_{timezone.now():%Y%m%d}.pdf"
        response = HttpResponse(buffer, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response